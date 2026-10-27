#!/usr/bin/env python3
"""Genera un libro nuevo del formato histórico para referencia y pruebas.

No rellena una plantilla del usuario. Consume datos ya preparados; no ejecuta
el proceso independiente de homologación ni valida decisiones técnicas.

Uso:
    python3 generar_cuadro_referencia.py entrada.json [carpeta_salida]

Orden de hojas que produce:
    1. CC_<ESP>-<PROY>   Hoja 1: formato original. Valores ESCRITOS, solo fórmulas internas,
                         comentarios cortos, colores de texto azul/rojo/negro.
    2. CC_<ESP>_<PROY>   Hoja 2: cuadro desglosado con leyenda de colores, columna
                         OBSERVACIONES / MOTIVO DEL RESALTADO, ratios, nivelación,
                         criterios y recomendación.
    3. LEYENDA Y NOTAS   Código de colores, resultado vinculado y conclusiones.
    4+. Una hoja por cotización (transcrita o copiada del Excel del postor).

El formato de entrada está en references/formato-entrada.md y hay ejemplos en assets/.
"""
import json
import math
import os
import re
import sys
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.colors import Color
from openpyxl.utils import get_column_letter as L

# ---------------------------------------------------------------- paleta
NAVY = "FF002060"
LIME = "FFCCFF33"
GREEN = "FF00B050"
ORANGE = "FFFFC000"
YELLOW = "FFFFFF00"
GOLD_LIGHT = "FFFFF2CC"
CELESTE = "FFCCFFFF"
YELLOW_LIGHT = "FFFFFFCC"
GRAY = "FFF2F2F2"
CH1 = "FF9DC3E6"
CH2 = "FFDEEBF7"
BLACK = "FF000000"
WHITE = "FFFFFFFF"
BLUE_TXT = "FF0070C0"
RED_TXT = "FFC00000"

FMT_MONEY = '_-"{m} "* #,##0.00_-;"-{m} "* #,##0.00_-;_-"{m} "* \\-??_-;_-@_-'
FMT_NUM = '_-* #,##0.00_-;\\-* #,##0.00_-;_-* \\-??_-;_-@_-'
FMT_QTY = '###,###,###,##0.00'
FMT_PCT = '0.0%'
FMT_SIMPLE = '"{m} "#,##0.00;"-{m} "#,##0.00'

TONO = {"ventaja": (BLUE_TXT, True), "riesgo": (RED_TXT, True),
        "faltante": (BLACK, True), "info": (BLACK, False)}

ESTADO_H2 = {  # relleno, color de letra del encabezado del postor en la Hoja 2
    "recomendado": (GREEN, WHITE), "mejor": (LIME, BLACK), "negociacion": (YELLOW, BLACK),
    "caro": (ORANGE, RED_TXT), "anterior": (BLACK, YELLOW), "normal": (LIME, BLACK)}

CRITERIOS_BASE = [
    "ADELANTO / TIPO DE GARANTIA", "FRECUENCIA DE VALORIZACIONES", "FONDO DE GARANTIA",
    "CONTACTO", "CONSIDERACIONES:", "N° DE OBRAS ACTUALES EN EJECUCION",
    "VALIDEZ DE LA OFERTA", "PLAZO DE EJECUCION", "GARANTIA"]

MESES = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO",
         "SETIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]


def fill(c):
    if isinstance(c, Color):
        return PatternFill(fill_type="solid", fgColor=c)
    return PatternFill(fill_type="solid", fgColor=c)


def theme(t, tint):
    return fill(Color(theme=t, tint=tint))


def font(sz=14, b=False, color=None, i=False):
    return Font(name="Arial", size=sz, bold=b, italic=i, color=color)


THIN = Side(style="thin", color="FFBFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_TL = Alignment(wrap_text=True, vertical="top", horizontal="left")
WRAP_CC = Alignment(wrap_text=True, vertical="center", horizontal="center")
WRAP_CL = Alignment(wrap_text=True, vertical="center", horizontal="left")


def num(v):
    """Valor numérico de un número o de una fórmula interna simple ('=50900/2000')."""
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().lstrip("=")
    if re.fullmatch(r"[0-9.+\-*/() ]+", s):
        return float(eval(s, {"__builtins__": {}}))  # solo aritmética
    raise ValueError(f"Valor no numérico: {v!r}")


def comment(cell, text, w=320, h=110):
    c = Comment(text, "Costos")
    c.width, c.height = w, h
    cell.comment = c


def lines_needed(text, width_chars):
    n = 0
    for part in str(text).split("\n"):
        n += max(1, math.ceil(len(part) / max(width_chars, 1)))
    return n


def ranges(rows):
    """[13,14,15,18] -> 'X13:X15,X18' (se completa con la columna afuera)."""
    rows = sorted(rows)
    out, start, prev = [], None, None
    for r in rows:
        if start is None:
            start = prev = r
        elif r == prev + 1:
            prev = r
        else:
            out.append((start, prev))
            start = prev = r
    if start is not None:
        out.append((start, prev))
    return out


def sum_formula(col, rows):
    if not rows:
        return 0
    parts = [f"{col}{a}" if a == b else f"{col}{a}:{col}{b}" for a, b in ranges(rows)]
    return f"=SUM({','.join(parts)})"


def safe_sheet(name, used):
    name = re.sub(r"[\[\]:*?/\\]", "-", name)[:31].strip("'")
    base, k = name, 2
    while name.upper() in used:
        suf = f" ({k})"
        name = base[:31 - len(suf)] + suf
        k += 1
    used.add(name.upper())
    return name


def fecha_larga(f):
    d = datetime.strptime(f, "%Y-%m-%d").date() if isinstance(f, str) else (f or date.today())
    return d, f"FECHA: {d.day} DE {MESES[d.month - 1]} DEL {d.year}"


# ---------------------------------------------------------------- modelo
class Modelo:
    def __init__(self, spec):
        self.s = spec
        self.moneda = spec.get("moneda", "S/")
        self.usd = bool(spec.get("tipo_cambio"))
        self.tc = float(spec.get("tipo_cambio") or 1)
        self.igv = float(spec.get("igv", 0.18))
        self.umbral = float(spec.get("umbral_alto_impacto", 10000))
        self.postores = spec["postores"]
        self.ids = [p["id"] for p in self.postores]
        self.P = {p["id"]: p for p in self.postores}
        self.base = spec.get("base_homologacion") or self.ids[0]
        self.fmt_m = FMT_MONEY.format(m="S/")
        self.fmt_usd = FMT_MONEY.format(m="US$")
        self.fmt_s = FMT_SIMPLE.format(m="S/")
        self.items = []  # (item, ancestros)
        self._walk(spec["partidas"], [])
        self.calc()

    def desc(self, pid):
        d = self.P[pid].get("descuento")
        return None if d is None else float(d)

    def _walk(self, nodes, anc):
        for n in nodes:
            if "hijos" in n:
                self._walk(n["hijos"], anc + [n])
            else:
                for it in n.get("items", []):
                    self.items.append((it, anc + [n]))

    def global_de(self, anc, pid):
        for ch in anc:
            b = (ch.get("bloques") or {}).get(pid)
            if b and b.get("cant") is not None:
                return ch, b
        return None, None

    def bloque_de(self, anc, pid):
        for ch in anc:
            b = (ch.get("bloques") or {}).get(pid)
            if b:
                return ch, b
        return None, None

    def precio(self, it, pid):
        """Devuelve (cant, pu_soles, parcial_neto) o None."""
        p = (it.get("precios") or {}).get(pid)
        if not p or p.get("cant") is None:
            return None
        cant = num(p["cant"])
        pu = num(p.get("pu_usd")) * self.tc if p.get("pu_usd") is not None else num(p.get("pu"))
        d = self.desc(pid) or 0
        return cant, pu, cant * pu * (1 - d)

    def cubierto(self, it, anc, pid):
        if self.precio(it, pid):
            return "precio"
        if self.global_de(anc, pid)[0] is not None:
            return "global"
        if (it.get("precios") or {}).get(pid, {}).get("incluido"):
            return "incluido"
        return None

    def clase(self, it, anc):
        if it.get("clase"):
            return it["clase"]
        priced = [pid for pid in self.ids if self.precio(it, pid)]
        cov = [pid for pid in self.ids if self.cubierto(it, anc, pid)]
        if self.base not in cov and len(priced) == 1:
            return "solo_uno"
        if self.base not in cov and len(priced) >= 2:
            return "varios_fuera_base"
        if len(cov) < len(self.ids) and priced:
            return "omitida"
        parc = [self.precio(it, pid)[2] for pid in priced]
        if len(parc) >= 2 and max(parc) - min(parc) > self.umbral:
            return "alto_impacto"
        return "normal"

    def calc(self):
        self.tot = {}
        for pid in self.ids:
            t = 0.0
            seen = set()
            for it, anc in self.items:
                ch, b = self.global_de(anc, pid)
                if ch is not None:
                    if id(ch) not in seen:
                        seen.add(id(ch))
                        pu = num(b.get("pu_usd")) * self.tc if b.get("pu_usd") is not None else num(b["pu"])
                        t += num(b["cant"]) * pu * (1 - (self.desc(pid) or 0))
                    continue
                pr = self.precio(it, pid)
                if pr and not it["precios"][pid].get("excluir_subtotal"):
                    t += pr[2]
            p = self.P[pid]
            t += float(p.get("ajuste_declarado") or 0)
            cd = t
            sub = cd * (1 + float(p.get("gg") or 0) + float(p.get("utilidad") or 0))
            self.tot[pid] = {"cd": cd, "sub": sub, "total": sub * (1 + self.igv)}


# ---------------------------------------------------------------- columnas
def bloques_columnas(m, start, hoja):
    """Mapa de columnas por postor. hoja=1 o 2."""
    cols, c = {}, start
    for pid in m.ids:
        d = {"cant": c}
        c += 1
        if m.usd:
            d["pu_usd"] = c
            c += 1
        d["pu"] = c
        d["parcial"] = c + 1
        c += 2
        if m.desc(pid) is not None:
            d["pu_neto"], d["parcial_neto"] = c, c + 1
            c += 2
        d["tot"] = d.get("parcial_neto", d["parcial"])
        d["ini"], d["fin"] = d["cant"], c - 1
        cols[pid] = d
    return cols, c


def encabezados(ws, m, cols, sz, h2):
    for pid, d in cols.items():
        p = m.P[pid]
        ws.merge_cells(start_row=6, start_column=d["ini"], end_row=6, end_column=d["fin"])
        top = ws.cell(6, d["ini"], p["nombre"])
        bg, fg = ESTADO_H2.get(p.get("estado", "normal"), (LIME, BLACK)) if h2 else (LIME, BLACK)
        for c in range(d["ini"], d["fin"] + 1):
            ws.cell(6, c).fill = fill(bg)
            ws.cell(7, c).fill = fill(bg if h2 else LIME)
        top.font = font(sz, True, fg)
        top.alignment = WRAP_CC
        mon = "S/"
        heads = {"cant": "Cant.", "pu_usd": "P.Unitario (US$)", "pu": f"P.Unitario ({mon})",
                 "parcial": f"P.Parcial ({mon})"}
        dsc = m.desc(pid)
        if dsc is not None:
            heads["pu_neto"] = f"P.Unitario ({mon}) INC DSCTO {dsc * 100:.2f}%"
            heads["parcial_neto"] = f"P.Parcial ({mon}) INC DSCTO {dsc * 100:.2f}%"
        for k, t in heads.items():
            if k in d:
                x = ws.cell(7, d[k], t)
                x.font = font(sz, True, fg if h2 else None)
                x.alignment = WRAP_CC


def meta_heads(ws, sz):
    ws.merge_cells("B6:F6")
    ws["B6"] = "PRESUPUESTO META"
    for c, t in zip("BCDEF", ["Descripción", "Unid", "Metrado ", "P.Unitario\n (S/)", "P.Parcial\n (S/)"]):
        ws[f"{c}6"].fill = fill(NAVY)
        x = ws[f"{c}7"]
        x.value = t
        x.fill = fill(NAVY)
        x.font = font(sz, True, WHITE)
        x.alignment = WRAP_CC
    ws["B6"].font = font(sz, True, WHITE)
    ws["B6"].alignment = WRAP_CC


def titulos(ws, m, sz_base):
    s = m.s
    _, f = fecha_larga(s["proyecto"].get("fecha"))
    ws["B2"] = s["proyecto"]["obra"]
    ws["B3"] = s["especialidad"]["titulo"]
    ws["B4"] = f
    ws["B2"].font = font(sz_base + 10, True)
    ws["B3"].font = font(sz_base + 6, True)
    ws["B4"].font = font(sz_base, True)
    if m.usd:
        ws["E5"] = "TC"
        ws["F5"] = m.tc
        ws["E5"].font = font(sz_base, True)
        ws["F5"].font = font(sz_base, True, BLUE_TXT)
        comment(ws["F5"], "Tipo de cambio S/ por US$ (editable).")


def write_price(ws, m, r, d, pid, cant, pu, pu_usd, f_txt, money_fmt, rate_cell):
    ws.cell(r, d["cant"], cant).number_format = FMT_QTY
    if m.usd and pu_usd is not None:
        ws.cell(r, d["pu_usd"], pu_usd).number_format = m.fmt_usd
        ws.cell(r, d["pu"], f"={L(d['pu_usd'])}{r}*$F$5").number_format = money_fmt
    else:
        ws.cell(r, d["pu"], pu).number_format = money_fmt
    ws.cell(r, d["parcial"], f"={L(d['pu'])}{r}*{L(d['cant'])}{r}").number_format = money_fmt
    if "pu_neto" in d:
        ws.cell(r, d["pu_neto"], f"={L(d['pu'])}{r}*(1-{rate_cell})").number_format = money_fmt
        ws.cell(r, d["parcial_neto"], f"={L(d['pu_neto'])}{r}*{L(d['cant'])}{r}").number_format = money_fmt
    for k in d:
        if k in ("cant", "pu_usd", "pu", "parcial", "pu_neto", "parcial_neto"):
            ws.cell(r, d[k]).font = f_txt


# ---------------------------------------------------------------- hoja 1
def hoja1(wb, m, name):
    ws = wb.active
    ws.title = name
    sz = 14
    ws.sheet_view.zoomScale = 60
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = NAVY[2:]
    cols, end = bloques_columnas(m, 8, 1)
    widths = {"A": 3.3, "B": 107.6, "C": 7.6, "D": 13.2, "E": 17.1, "F": 29.3, "G": 4.1}
    for k, v in widths.items():
        ws.column_dimensions[k].width = v
    for d in cols.values():
        for k, w in (("cant", 13), ("pu_usd", 20), ("pu", 22), ("parcial", 34), ("pu_neto", 22), ("parcial_neto", 34)):
            if k in d:
                ws.column_dimensions[L(d[k])].width = w
    titulos(ws, m, sz)
    meta_heads(ws, sz)
    encabezados(ws, m, cols, sz, False)
    ws.row_dimensions[7].height = 35.25
    for r in range(2, 8):
        for c in range(2, end):
            if ws.cell(r, c).font.name != "Arial":
                ws.cell(r, c).font = font(sz, ws.cell(r, c).font.b)
    ws.freeze_panes = "G8"

    rate_cells = {}
    r = 9
    total_row = r
    titulo_total = m.s["especialidad"].get("titulo_total", m.s["especialidad"]["nombre"].upper())
    ws.cell(r, 2, titulo_total)
    _paint_row(ws, r, end, theme(8, 0.6), font(sz, True))
    for pid, d in cols.items():
        if "pu_neto" in d:
            x = ws.cell(r, d["pu_neto"], m.desc(pid))
            x.number_format = "0.00%"
            x.font = font(sz, True, BLUE_TXT)
            rate_cells[pid] = f"${L(d['pu_neto'])}${r}"
    r += 2
    ctx = {"ws": ws, "m": m, "cols": cols, "end": end, "sz": sz, "rates": rate_cells,
           "hay_meta": any(it.get("meta") for it, _ in m.items)}
    child_rows, r = _h1_nodos(ctx, m.s["partidas"], r, 1)
    for c in ["F"] + [L(d["tot"]) for d in cols.values()]:
        ws[f"{c}{total_row}"] = sum_formula(c, child_rows)
        ws[f"{c}{total_row}"].number_format = m.fmt_m
        ws[f"{c}{total_row}"].font = font(sz, True)
    for pid, d in cols.items():
        if d["tot"] != d["parcial"]:
            ws.cell(total_row, d["parcial"], sum_formula(L(d["parcial"]), child_rows)).number_format = m.fmt_m

    # ajustes declarados
    if any(p.get("ajuste_declarado") for p in m.postores):
        r += 1
        ws.cell(r, 2, "AJUSTE A PRECIO DECLARADO").font = font(sz, True)
        for pid, d in cols.items():
            aj = m.P[pid].get("ajuste_declarado")
            if aj:
                x = ws.cell(r, d["tot"], float(aj))
                x.number_format = m.fmt_m
                x.font = font(sz, True, RED_TXT)
                comment(x, "Diferencia entre el total declarado y la suma de sus partidas. Se respeta el precio declarado.")
        adj_row = r
        for pid, d in cols.items():
            col = L(d["tot"])
            ws[f"{col}{total_row}"] = f"={ws[f'{col}{total_row}'].value.lstrip('=')}+{col}{adj_row}"
        r += 1

    # pie
    r += 3
    pie = r
    labels = ["PARCIAL (SIN IGV)", "GASTOS GENERALES", "UTILIDADES", "SUBTOTAL (SIN IGV)",
              f"IGV  ({int(m.igv * 100)}%)", "TOTAL INC. IGV"]
    for i, t in enumerate(labels):
        rr = pie + i
        ws.merge_cells(f"D{rr}:E{rr}")
        ws[f"D{rr}"].fill = fill(NAVY)
        for c, v in (("D", t), ("E", None), ("F", None)):
            x = ws[f"{c}{rr}"]
            if v:
                x.value = v
            x.fill = fill(NAVY)
            x.font = font(sz, True, WHITE)
            x.alignment = Alignment(horizontal="right", vertical="center")
    ws[f"F{pie}"] = f"=F{total_row}"
    ws[f"F{pie + 1}"] = 0
    ws[f"F{pie + 2}"] = 0
    ws[f"F{pie + 3}"] = f"=SUM(F{pie}:F{pie + 2})"
    ws[f"F{pie + 4}"] = f"=+F{pie + 3}*{m.igv}"
    ws[f"F{pie + 5}"] = f"=+F{pie + 3}+F{pie + 4}"
    totales = {pid: m.tot[pid]["total"] for pid in m.ids}
    tmin, tmax = min(totales.values()), max(totales.values())
    for pid, d in cols.items():
        col = L(d["tot"])
        p = m.P[pid]
        ws[f"{col}{pie}"] = f"={col}{total_row}"
        ws[f"{col}{pie + 1}"] = f"={col}{pie}*{float(p['gg'])}" if p.get("gg") else 0
        ws[f"{col}{pie + 2}"] = f"={col}{pie}*{float(p['utilidad'])}" if p.get("utilidad") else 0
        ws[f"{col}{pie + 3}"] = f"=SUM({col}{pie}:{col}{pie + 2})"
        ws[f"{col}{pie + 4}"] = f"=+{col}{pie + 3}*{m.igv}"
        ws[f"{col}{pie + 5}"] = f"=+{col}{pie + 3}+{col}{pie + 4}"
        for i in range(6):
            ws[f"{col}{pie + i}"].font = font(sz, True)
        if len(m.ids) > 1:
            if totales[pid] == tmin:
                ws[f"{col}{pie + 5}"].font = font(sz, True, BLUE_TXT)
            elif totales[pid] == tmax:
                ws[f"{col}{pie + 5}"].font = font(sz, True, RED_TXT)
        if p.get("gg") or p.get("utilidad"):
            comment(ws[f"{col}{pie + 1}"], f"GG {float(p.get('gg') or 0) * 100:.0f}% y utilidad "
                    f"{float(p.get('utilidad') or 0) * 100:.0f}% sobre el costo directo, según su cotización.")
    for c in ["F"] + [L(d["tot"]) for d in cols.values()]:
        for i in range(6):
            ws[f"{c}{pie + i}"].number_format = m.fmt_m
    tot_row = pie + 5
    totals = ",".join(f"${L(d['tot'])}${tot_row}" for d in cols.values())

    # DELTA / VARIACION / ratios
    r = tot_row + 2
    filas = [("DELTA", lambda d: f'=IF($F${tot_row}=0,"Ingresar meta",$F${tot_row}-{L(d["tot"])}{tot_row})', m.fmt_s),
             ("VARIACIÓN vs. OFERTA MÁS BAJA", lambda d: f"={L(d['tot'])}{tot_row}/MIN({totals})-1", FMT_PCT)]
    for ra in m.s.get("ratios", []):
        filas.append((ra["nombre"], lambda d, ra=ra: f"={L(d['tot'])}{pie + 3}/{ra['divisor']}", m.fmt_s))
    for lab, fx, nf in filas:
        ws.merge_cells(f"B{r}:F{r}")
        ws[f"B{r}"] = lab
        ws[f"B{r}"].fill = fill(NAVY)
        ws[f"B{r}"].font = font(sz, True, WHITE)
        ws[f"B{r}"].alignment = Alignment(horizontal="right", vertical="center")
        for pid, d in cols.items():
            ws.merge_cells(start_row=r, start_column=d["ini"], end_row=r, end_column=d["fin"])
            x = ws.cell(r, d["ini"], fx(d))
            x.number_format = nf
            x.font = font(sz, True)
            x.alignment = WRAP_CC
            for c in range(d["ini"], d["fin"] + 1):
                ws.cell(r, c).fill = fill(LIME)
        ws.row_dimensions[r].height = 20.25
        r += 1

    # criterios
    crit = m.s.get("criterios_hoja1") or (CRITERIOS_BASE[:] + m.s.get("criterios_extra", []))
    for lab in crit:
        es_cons = lab.upper().startswith("CONSIDERACIONES")
        ws.cell(r, 6, lab).font = font(12, False, None)
        ws.cell(r, 6).fill = theme(0, -0.05)
        ws.cell(r, 6).alignment = WRAP_CL
        maxl = 1
        for pid, d in cols.items():
            p = m.P[pid]
            if es_cons:
                txt = "\n".join(f"- {t}" for t in p.get("consideraciones", [])) or "No indica."
                tono = "info"
                sz_t = 11
            elif lab.upper().startswith("CONTACTO"):
                txt = p.get("contacto") or "No indica."
                tono, sz_t = "info", 11
            else:
                cr = (p.get("criterios") or {}).get(lab) or {}
                txt = cr.get("texto", "No indica.") if isinstance(cr, dict) else str(cr)
                tono = cr.get("tono", "faltante" if txt == "No indica." else "info") if isinstance(cr, dict) else "info"
                sz_t = 12
            ws.merge_cells(start_row=r, start_column=d["ini"], end_row=r, end_column=d["fin"])
            x = ws.cell(r, d["ini"], txt)
            colr, b = TONO.get(tono, TONO["info"])
            x.font = font(sz_t, b, colr)
            x.alignment = WRAP_TL if sz_t == 11 else WRAP_CC
            for c in range(d["ini"], d["fin"] + 1):
                ws.cell(r, c).border = BOX
            width = sum(ws.column_dimensions[L(c)].width for c in range(d["ini"], d["fin"] + 1))
            maxl = max(maxl, lines_needed(txt, width * 1.05))
        ws.cell(r, 6).border = BOX
        ws.row_dimensions[r].height = max(20.25, maxl * 15.5 + 6)
        r += 1
    return ws, cols, pie


def _paint_row(ws, r, end, f, fnt):
    for c in range(2, end):
        if c == 7:
            continue
        ws.cell(r, c).fill = f
        ws.cell(r, c).font = fnt


H1_FILLS = {1: (8, 0.8), 2: (0, -0.15), 3: (9, 0.8)}


def _h1_nodos(ctx, nodes, r, nivel):
    ws, m, cols, end, sz = ctx["ws"], ctx["m"], ctx["cols"], ctx["end"], ctx["sz"]
    rows_nivel = []
    for n in nodes:
        ch_row = r
        rows_nivel.append(ch_row)
        ws.cell(r, 2, n["titulo"])
        t, tint = H1_FILLS.get(nivel, (9, 0.8))
        _paint_row(ws, r, end, theme(t, tint), font(sz, True))
        if n.get("und"):
            ws.cell(r, 3, n["und"]).font = font(sz)
        meta = n.get("meta")
        if meta:
            ws.cell(r, 4, meta.get("metrado")).number_format = FMT_QTY
            ws.cell(r, 5, meta.get("pu")).number_format = FMT_NUM
        elif nivel == 1 and not ctx["hay_meta"]:
            e = ws.cell(r, 5)
            e.fill = fill(GOLD_LIGHT)
            comment(e, "Presupuesto meta: dato no incluido en las cotizaciones. Ingresar para calcular el DELTA.")
        ws.row_dimensions[r].height = 17.4
        r += 1
        ws.row_dimensions[r].height = 6
        r += 1
        first_child = r
        if "hijos" in n:
            child_rows, r = _h1_nodos(ctx, n["hijos"], r, nivel + 1)
        else:
            child_rows, r = _h1_items(ctx, n.get("items", []), r)
        last_child = r - 1
        # subtotales del capítulo
        if meta:
            ws.cell(ch_row, 6, f"=E{ch_row}*D{ch_row}")
        else:
            ws.cell(ch_row, 6, sum_formula("F", child_rows))
        ws.cell(ch_row, 6).number_format = FMT_NUM
        ws.cell(ch_row, 6).font = font(sz, True)
        for pid, d in cols.items():
            b = (n.get("bloques") or {}).get(pid)
            if b and b.get("cant") is not None:
                pu_usd = b.get("pu_usd")
                pu = num(b.get("pu")) if pu_usd is None else None
                write_price(ws, m, ch_row, d, pid, num(b["cant"]), pu, pu_usd, font(sz), m.fmt_m, ctx["rates"].get(pid))
                for k in ("parcial", "parcial_neto"):
                    if k in d:
                        ws.cell(ch_row, d[k]).font = font(sz, True)
                if b.get("comentario"):
                    comment(ws.cell(ch_row, d["pu"]), b["comentario"])
            else:
                excl = set(ctx.get("excl", {}).get(pid, []))
                rows = [x for x in child_rows if x not in excl]
                for k in ("parcial", "parcial_neto"):
                    if k in d:
                        x = ws.cell(ch_row, d[k], sum_formula(L(d[k]), rows))
                        x.number_format = m.fmt_m
                        x.font = font(sz, True)
                if b and b.get("comentario_subtotal"):
                    comment(ws.cell(ch_row, d["tot"]), b["comentario_subtotal"])
            if b and b.get("texto") and last_child >= first_child:
                ws.merge_cells(start_row=first_child, start_column=d["ini"], end_row=last_child, end_column=d["fin"])
                x = ws.cell(first_child, d["ini"], b["texto"])
                x.font = font(11, True, RED_TXT)
                x.alignment = WRAP_TL
    return rows_nivel, r


def _h1_items(ctx, items, r):
    ws, m, cols, end, sz = ctx["ws"], ctx["m"], ctx["cols"], ctx["end"], ctx["sz"]
    rows = []
    for it in items:
        rows.append(r)
        priced = [pid for pid in m.ids if m.precio(it, pid)]
        b = ws.cell(r, 2, it["desc"])
        b.font = font(sz, False, BLUE_TXT if len(priced) >= 2 else None)
        ws.cell(r, 3, it.get("und")).font = font(sz)
        meta = it.get("meta") or {}
        ws.cell(r, 4, meta.get("metrado")).number_format = FMT_QTY
        ws.cell(r, 5, meta.get("pu")).number_format = FMT_NUM
        ws.cell(r, 6, f"=E{r}*D{r}").number_format = FMT_NUM
        for c in (4, 5, 6):
            ws.cell(r, c).font = font(sz)
        for pid, d in cols.items():
            p = (it.get("precios") or {}).get(pid) or {}
            if p.get("cant") is not None:
                obs = bool(p.get("observado"))
                f = font(sz, obs, RED_TXT if obs else None)
                pu = None if p.get("pu_usd") is not None else p.get("pu")
                if isinstance(pu, str) and pu.startswith("="):
                    pu_val = pu
                else:
                    pu_val = num(pu) if pu is not None else None
                write_price(ws, m, r, d, pid, num(p["cant"]), pu_val, p.get("pu_usd"), f, m.fmt_m, ctx["rates"].get(pid))
                if p.get("nota"):
                    comment(ws.cell(r, d["pu"]), p["nota"])
                if p.get("excluir_subtotal"):
                    ctx.setdefault("excl", {}).setdefault(pid, []).append(r)
            elif p.get("incluido"):
                x = ws.cell(r, d["cant"], p["incluido"])
                no = str(p["incluido"]).lower().startswith("no ")
                x.font = font(11, True, RED_TXT if no else BLACK)
        ws.row_dimensions[r].height = 17.4
        r += 1
    return rows, r


# ---------------------------------------------------------------- hoja 2
CLASE_FILL = {
    "solo_uno": (GOLD_LIGHT, RED_TXT, False),
    "varios_fuera_base": (GOLD_LIGHT, BLUE_TXT, False),
    "omitida": (CELESTE, None, False),
    "alto_impacto": (YELLOW_LIGHT, None, False),
    "alerta": (YELLOW, RED_TXT, True),
    "referencial": (GRAY, None, True),
    "global": (GRAY, None, True),
    "normal": (None, None, False),
}


def hoja2(wb, m, name, refs):
    ws = wb.create_sheet(name)
    sz = 12
    ws.sheet_view.zoomScale = 80
    ws.sheet_view.showGridLines = False
    cols, end = bloques_columnas(m, 8, 2)
    obs_col = end + 1
    for k, v in {"A": 3, "B": 66, "C": 8, "D": 11, "E": 12, "F": 15, "G": 1.6}.items():
        ws.column_dimensions[k].width = v
    for d in cols.values():
        for k, w in (("cant", 11), ("pu_usd", 15), ("pu", 15), ("parcial", 18), ("pu_neto", 15), ("parcial_neto", 18)):
            if k in d:
                ws.column_dimensions[L(d[k])].width = w
    ws.column_dimensions[L(end)].width = 1.6
    ws.column_dimensions[L(obs_col)].width = 80
    titulos(ws, m, sz)
    ws["B2"].font = font(18, True)
    ws["B3"].font = font(14, True)
    meta_heads(ws, sz)
    encabezados(ws, m, cols, sz, True)
    x = ws.cell(6, obs_col, "OBSERVACIONES / MOTIVO DEL RESALTADO")
    ws.merge_cells(start_row=6, start_column=obs_col, end_row=7, end_column=obs_col)
    x.fill = fill(NAVY)
    x.font = font(sz, True, WHITE)
    x.alignment = WRAP_CC
    ws.row_dimensions[7].height = 45
    ws.freeze_panes = "G8"

    r = 9
    total_row = r
    ws.cell(r, 2, m.s["especialidad"].get("titulo_total", m.s["especialidad"]["nombre"].upper()))
    for c in range(2, end):
        if c != 7:
            ws.cell(r, c).fill = fill(BLACK)
            ws.cell(r, c).font = font(sz, True, ORANGE)
    rates = {}
    for pid, d in cols.items():
        if "pu_neto" in d:
            x = ws.cell(r, d["pu_neto"], m.desc(pid))
            x.number_format = "0.00%"
            x.font = font(sz, True, BLUE_TXT)
            x.fill = fill(WHITE)
            rates[pid] = f"${L(d['pu_neto'])}${r}"
            comment(x, "Descuento del postor (editable para simular negociación).")
    r += 2
    ctx = {"ws": ws, "m": m, "cols": cols, "end": end, "sz": sz, "rates": rates, "obs": obs_col,
           "refs": refs, "no_cot": {pid: [] for pid in m.ids}}
    child_rows, r = _h2_nodos(ctx, m.s["partidas"], r, 1)
    for c in ["F"] + [L(d["tot"]) for d in cols.values()]:
        ws[f"{c}{total_row}"] = sum_formula(c, child_rows)
    for d in cols.values():
        if d["tot"] != d["parcial"]:
            ws.cell(total_row, d["parcial"], sum_formula(L(d["parcial"]), child_rows))
    for c in range(2, end):
        if f"${L(c)}${total_row}" not in rates.values():
            ws.cell(total_row, c).number_format = m.fmt_m

    if any(p.get("ajuste_declarado") for p in m.postores):
        r += 1
        ws.cell(r, 2, "AJUSTE A PRECIO DECLARADO").font = font(sz, True)
        ws.cell(r, 2).fill = fill(YELLOW)
        ws.cell(r, 2).font = font(sz, True, RED_TXT)
        txt = []
        for pid, d in cols.items():
            aj = m.P[pid].get("ajuste_declarado")
            if aj:
                ws.cell(r, d["tot"], float(aj)).number_format = m.fmt_m
                txt.append(f"{m.P[pid]['corto']}: su detalle no suma su total declarado; se respeta el precio declarado "
                           f"(ajuste {m.moneda} {float(aj):,.2f}). Pedir cuadro corregido.")
        _obs(ws, r, obs_col, "ROJO SOBRE AMARILLO: ALERTA. " + " ".join(txt))
        for d in cols.values():
            col = L(d["tot"])
            ws[f"{col}{total_row}"] = f"={ws[f'{col}{total_row}'].value.lstrip('=')}+{col}{r}"
        r += 1

    # pie
    r += 2
    pie = r
    labels = ["COSTO DIRECTO (SIN IGV)", "GASTOS GENERALES", "UTILIDAD", "SUBTOTAL (SIN IGV)",
              f"IGV ({int(m.igv * 100)}%)", "TOTAL INC. IGV"]
    for i, t in enumerate(labels):
        ws.cell(pie + i, 2, t).font = font(sz, True, WHITE)
        ws.cell(pie + i, 2).fill = fill(NAVY)
    ws[f"F{pie}"] = f"=F{total_row}"
    ws[f"F{pie + 1}"] = 0
    ws[f"F{pie + 2}"] = 0
    ws[f"F{pie + 3}"] = f"=SUM(F{pie}:F{pie + 2})"
    ws[f"F{pie + 4}"] = f"=F{pie + 3}*{m.igv}"
    ws[f"F{pie + 5}"] = f"=F{pie + 3}+F{pie + 4}"
    ggtxt = []
    for pid, d in cols.items():
        p = m.P[pid]
        col, pc = L(d["tot"]), L(d["tot"] - 1)
        ws[f"{col}{pie}"] = f"={col}{total_row}"
        for i, k in ((1, "gg"), (2, "utilidad")):
            x = ws[f"{pc}{pie + i}"]
            x.value = float(p.get(k) or 0)
            x.number_format = "0.00%"
            x.font = font(sz, True, BLUE_TXT)
            ws[f"{col}{pie + i}"] = f"={pc}{pie + i}*{col}{pie}"
        ws[f"{col}{pie + 3}"] = f"=SUM({col}{pie}:{col}{pie + 2})"
        ws[f"{col}{pie + 4}"] = f"={col}{pie + 3}*{m.igv}"
        ws[f"{col}{pie + 5}"] = f"={col}{pie + 3}+{col}{pie + 4}"
        ggtxt.append(f"{p['corto']} {float(p.get('gg') or 0) * 100:.0f}% + {float(p.get('utilidad') or 0) * 100:.0f}%")
        for i in range(6):
            ws[f"{col}{pie + i}"].number_format = m.fmt_m
            ws[f"{col}{pie + i}"].font = font(sz, True)
    for i in range(6):
        ws[f"F{pie + i}"].number_format = m.fmt_m
    _obs(ws, pie + 1, obs_col, "GG + utilidad por postor (números azules editables): " + " · ".join(ggtxt) + ".")
    decl = []
    for pid in m.ids:
        dv = m.P[pid].get("total_declarado_sin_igv")
        if dv is not None:
            decl.append(f"{m.P[pid]['corto']} {m.moneda} {float(dv):,.2f}")
    if decl:
        _obs(ws, pie + 3, obs_col, "Debe cuadrar con cada cotización (sin IGV): " + " · ".join(decl) + ".")
    sub_row, tot_row = pie + 3, pie + 5
    totals = ",".join(f"${L(d['tot'])}${tot_row}" for d in cols.values())

    r = tot_row + 2
    for lab, fx, nf in (
            ("DELTA (PRESUPUESTO META - POSTOR, INC. IGV)",
             lambda d: f'=IF($F${tot_row}=0,"Falta ppto. meta",$F${tot_row}-{L(d["tot"])}{tot_row})', m.fmt_s),
            ("DIFERENCIA vs. MENOR OFERTA (INC. IGV)", lambda d: f"={L(d['tot'])}{tot_row}-MIN({totals})", m.fmt_s),
            ("VARIACIÓN % vs. MENOR OFERTA", lambda d: f"={L(d['tot'])}{tot_row}/MIN({totals})-1", FMT_PCT)):
        ws.cell(r, 2, lab).font = font(sz, True, WHITE)
        ws.cell(r, 2).fill = fill(NAVY)
        for d in cols.values():
            ws.merge_cells(start_row=r, start_column=d["ini"], end_row=r, end_column=d["fin"])
            x = ws.cell(r, d["ini"], fx(d))
            x.number_format = nf
            x.font = font(sz, True)
            x.alignment = WRAP_CC
            for c in range(d["ini"], d["fin"] + 1):
                ws.cell(r, c).fill = fill(LIME)
        r += 1

    # ratios
    r += 1
    _section(ws, m, cols, r, "ANÁLISIS DE PRECIOS UNITARIOS Y RATIOS (SIN IGV)", obs_col)
    r += 1
    for ra in m.s.get("ratios", []):
        _fila_valores(ws, cols, r, ra["nombre"], lambda d, ra=ra: f"={L(d['tot'])}{sub_row}/{ra['divisor']}", m.fmt_s,
                      obs_col, ra.get("obs"))
        r += 1
    for ch_row, titulo in ctx["capitulos"]:
        _fila_valores(ws, cols, r, f"Capítulo {titulo} ({m.moneda}, sin IGV)",
                      lambda d, ch_row=ch_row: f"={L(d['tot'])}{ch_row}", m.fmt_s, obs_col, None)
        r += 1
    for ra in m.s.get("ratios_extra", []):  # [{"nombre","valores":{pid: formula o valor},"obs"}]
        _fila_valores(ws, cols, r, ra["nombre"], lambda d, ra=ra, cols=cols: _valor_pid(ra, cols, d),
                      ra.get("formato", m.fmt_s), obs_col, ra.get("obs"))
        r += 1

    # nivelación
    r += 1
    _section(ws, m, cols, r, "NIVELACIÓN POR ALCANCE (SIN IGV)", obs_col)
    r += 1
    niv_add = r
    for pid, d in cols.items():
        terms = ctx["no_cot"][pid]
        ws.merge_cells(start_row=r, start_column=d["ini"], end_row=r, end_column=d["fin"])
        ws.cell(r, d["ini"], ("=" + "+".join(terms)) if terms else 0).number_format = m.fmt_s
    ws.cell(r, 2, "(+) Partidas no cotizadas (celestes y rojo/dorado), al menor precio de los otros postores").font = font(sz, True)
    _obs(ws, r, obs_col, "Escenario conservador: a cada postor se le suman las partidas que no cotiza ni declara "
                         "incluidas, al menor precio neto de los demás postores.")
    r += 1
    niv = r
    _fila_valores(ws, cols, r, "PRECIO NIVELADO (SIN IGV)", lambda d: f"={L(d['tot'])}{sub_row}+{L(d['ini'])}{niv_add}",
                  m.fmt_s, obs_col, None)
    r += 1
    _fila_valores(ws, cols, r, "PRECIO NIVELADO INC. IGV", lambda d: f"={L(d['ini'])}{niv}*{1 + m.igv}", m.fmt_s, obs_col, None)
    r += 1
    nivs = ",".join(f"${L(d['ini'])}${niv}" for d in cols.values())
    _fila_valores(ws, cols, r, "VARIACIÓN % NIVELADA vs. MENOR", lambda d: f"={L(d['ini'])}{niv}/MIN({nivs})-1", FMT_PCT,
                  obs_col, None)
    niv_var = r
    r += 1

    # secciones extra (flujo de pagos, comparativo técnico, etc.)
    for sec in m.s.get("secciones_extra", []):
        r += 1
        _section(ws, m, cols, r, sec["titulo"], obs_col)
        r += 1
        for fila in sec["filas"]:
            _fila_valores(ws, cols, r, fila["concepto"], lambda d, fila=fila, cols=cols: _valor_pid(fila, cols, d),
                          fila.get("formato", m.fmt_s if not sec.get("texto") else "General"), obs_col, fila.get("obs"),
                          texto=sec.get("texto", False), tonos=fila.get("tonos"))
            r += 1

    # criterios
    r += 1
    _section(ws, m, cols, r, "CRITERIOS COMERCIALES Y TÉCNICOS", obs_col)
    _obs(ws, r, obs_col, "Color del texto: AZUL negrita = ventaja / cumple · ROJO negrita = desventaja / riesgo de costo · "
                         "NEGRO negrita = dato faltante o a confirmar · negro normal = informativo")
    ws.cell(r, obs_col).font = font(sz, True)
    r += 1
    crit = ["EMPRESA / RUC / CONTACTO", "FECHA / N° / VERSIÓN", "MODALIDAD / ALCANCE", "GG / UTILIDAD / DESCUENTO"]
    crit += [c for c in (m.s.get("criterios_hoja1") or CRITERIOS_BASE + m.s.get("criterios_extra", []))
             if not c.upper().startswith(("CONTACTO", "CONSIDERACIONES"))]
    crit += m.s.get("criterios_hoja2_extra", []) + ["CONSIDERACIONES Y EXCLUSIONES"]
    for lab in crit:
        ws.cell(r, 2, lab).font = font(sz, True)
        ws.cell(r, 2).alignment = WRAP_CL
        maxl = 1
        for pid, d in cols.items():
            p = m.P[pid]
            if lab == "EMPRESA / RUC / CONTACTO":
                txt = " - ".join(x for x in [p["nombre"], ("RUC " + p["ruc"]) if p.get("ruc") else "RUC: por confirmar"] if x)
                if p.get("contacto"):
                    txt += "\n" + p["contacto"].replace("\n", " - ")
                tono = "info" if p.get("ruc") else "faltante"
            elif lab == "FECHA / N° / VERSIÓN":
                txt = p.get("referencia_cotizacion", "No indica.")
                tono = "info"
            elif lab == "MODALIDAD / ALCANCE":
                cr = p.get("modalidad") or {}
                txt, tono = (cr.get("texto", "No indica."), cr.get("tono", "info")) if isinstance(cr, dict) else (str(cr), "info")
            elif lab == "GG / UTILIDAD / DESCUENTO":
                txt = (f"GG {float(p.get('gg') or 0) * 100:.0f}% · Utilidad {float(p.get('utilidad') or 0) * 100:.0f}% · "
                       f"Descuento {float(p.get('descuento') or 0) * 100:.2f}%")
                tono = "info"
            elif lab == "CONSIDERACIONES Y EXCLUSIONES":
                txt = "\n".join(f"- {t}" for t in p.get("consideraciones", [])) or "No indica."
                tono = "info"
            else:
                cr = (p.get("criterios") or {}).get(lab) or {}
                txt = cr.get("texto", "No indica.") if isinstance(cr, dict) else str(cr)
                tono = cr.get("tono", "faltante" if txt == "No indica." else "info") if isinstance(cr, dict) else "info"
            ws.merge_cells(start_row=r, start_column=d["ini"], end_row=r, end_column=d["fin"])
            x = ws.cell(r, d["ini"], txt)
            colr, b = TONO.get(tono, TONO["info"])
            x.font = font(10, b, colr)
            x.alignment = WRAP_TL
            width = sum(ws.column_dimensions[L(c)].width for c in range(d["ini"], d["fin"] + 1))
            maxl = max(maxl, lines_needed(txt, width * 1.25))
        ws.row_dimensions[r].height = max(18, maxl * 13 + 4)
        r += 1

    # recomendación
    r += 1
    ws.cell(r, 2, "RECOMENDACIÓN").font = font(sz, True, WHITE)
    ws.cell(r, 2).fill = fill(NAVY)
    first, last = cols[m.ids[0]]["ini"], obs_col
    ws.merge_cells(start_row=r, start_column=first, end_row=r, end_column=last)
    rec = m.s.get("recomendacion", "Pendiente.")
    x = ws.cell(r, first, rec)
    x.font = font(sz, True)
    x.alignment = WRAP_TL
    width = sum(ws.column_dimensions[L(c)].width for c in range(first, last + 1))
    ws.row_dimensions[r].height = max(30, lines_needed(rec, width * 1.0) * 16 + 6)
    return ws, cols, {"sub": sub_row, "niv": niv, "niv_var": niv_var, "tot": tot_row}


def _valor_pid(fila, cols, d):
    for pid, dd in cols.items():
        if dd is d:
            return (fila.get("valores") or {}).get(pid, "")
    return ""


def _section(ws, m, cols, r, titulo, obs_col):
    ws.cell(r, 2, titulo).font = font(12, True, WHITE)
    ws.cell(r, 2).fill = fill(NAVY)
    for pid, d in cols.items():
        ws.merge_cells(start_row=r, start_column=d["ini"], end_row=r, end_column=d["fin"])
        x = ws.cell(r, d["ini"], m.P[pid]["corto"])
        x.font = font(12, True, WHITE)
        x.alignment = WRAP_CC
        for c in range(d["ini"], d["fin"] + 1):
            ws.cell(r, c).fill = fill(NAVY)


def _fila_valores(ws, cols, r, label, fx, nf, obs_col, obs, texto=False, tonos=None):
    ws.cell(r, 2, label).font = font(12, True)
    ws.cell(r, 2).alignment = WRAP_CL
    maxl = 1
    for pid, d in cols.items():
        ws.merge_cells(start_row=r, start_column=d["ini"], end_row=r, end_column=d["fin"])
        v = fx(d)
        x = ws.cell(r, d["ini"], v)
        x.number_format = nf
        tono = (tonos or {}).get(pid, "info")
        colr, b = TONO.get(tono, TONO["info"])
        x.font = font(10 if texto else 12, b if texto else True, colr)
        x.alignment = WRAP_TL if texto else WRAP_CC
        if texto:
            width = sum(ws.column_dimensions[L(c)].width for c in range(d["ini"], d["fin"] + 1))
            maxl = max(maxl, lines_needed(str(v), width * 1.25))
    if texto:
        ws.row_dimensions[r].height = max(18, maxl * 13 + 4)
    if obs:
        _obs(ws, r, obs_col, obs)


def _obs(ws, r, obs_col, text):
    x = ws.cell(r, obs_col, text)
    x.font = font(10)
    x.alignment = WRAP_TL
    need = lines_needed(text, 80 * 1.2) * 13 + 4
    cur = ws.row_dimensions[r].height or 15
    ws.row_dimensions[r].height = max(cur, need)


H2_FILLS = {1: CH1, 2: CH2, 3: GRAY}


def _h2_nodos(ctx, nodes, r, nivel):
    ws, m, cols, end, sz = ctx["ws"], ctx["m"], ctx["cols"], ctx["end"], ctx["sz"]
    ctx.setdefault("capitulos", [])
    rows_nivel = []
    for n in nodes:
        ch_row = r
        rows_nivel.append(r)
        if nivel == 1:
            ctx["capitulos"].append((r, n["titulo"]))
        ws.cell(r, 2, n["titulo"])
        for c in range(2, end):
            if c != 7:
                ws.cell(r, c).fill = fill(H2_FILLS.get(nivel, GRAY))
                ws.cell(r, c).font = font(sz, True)
        if n.get("meta"):
            ws.cell(r, 3, n.get("und"))
            ws.cell(r, 4, n["meta"].get("metrado"))
            ws.cell(r, 5, n["meta"].get("pu"))
        r += 1
        ctx.setdefault("anc", []).append(n)
        child_rows = []
        # fila gris por postor con precio global en este capítulo
        for pid, b in (n.get("bloques") or {}).items():
            if b.get("cant") is None:
                continue
            d = cols[pid]
            child_rows.append(r)
            ws.cell(r, 2, f"{n['titulo']} - PRECIO GLOBAL {m.P[pid]['corto'].upper()}")
            ws.cell(r, 3, n.get("und") or "glb")
            ws.cell(r, 6, f"=PRODUCT(D{r},E{r})").number_format = m.fmt_m
            for c in range(2, 7):
                ws.cell(r, c).fill = fill(GRAY)
                ws.cell(r, c).font = font(sz, True)
            pu_usd = b.get("pu_usd")
            write_price(ws, m, r, d, pid, num(b["cant"]), None if pu_usd is not None else num(b["pu"]), pu_usd,
                        font(sz), m.fmt_m, ctx["rates"].get(pid))
            _obs(ws, r, ctx["obs"], f"GRIS: Partida GLOBAL de {m.P[pid]['corto']}. {b.get('texto', '')} "
                                    f"Comparar por SUBTOTAL del capítulo.".replace("  ", " "))
            r += 1
        if "hijos" in n:
            rows, r = _h2_nodos(ctx, n["hijos"], r, nivel + 1)
        else:
            rows, r = _h2_items(ctx, n, r)
        child_rows += rows
        ctx["anc"].pop()
        ws.cell(ch_row, 6, (f"=E{ch_row}*D{ch_row}" if n.get("meta") else sum_formula("F", child_rows)))
        for d in cols.values():
            for k in ("parcial", "parcial_neto"):
                if k in d:
                    ws.cell(ch_row, d[k], sum_formula(L(d[k]), child_rows))
        for c in range(4, end):
            ws.cell(ch_row, c).number_format = m.fmt_m
    return rows_nivel, r


def _h2_items(ctx, n, r):
    ws, m, cols, sz, obs_col = ctx["ws"], ctx["m"], ctx["cols"], ctx["sz"], ctx["obs"]
    rows = []
    anc = list(ctx["anc"])
    anc_global = {pid for pid in m.ids if m.global_de(anc, pid)[0] is not None}
    for it in n.get("items", []):
        rows.append(r)
        clase = it.get("clase") or m.clase(it, anc)
        ws.cell(r, 2, it["desc"])
        ws.cell(r, 3, it.get("und"))
        meta = it.get("meta") or {}
        ws.cell(r, 4, meta.get("metrado"))
        ws.cell(r, 5, meta.get("pu"))
        ws.cell(r, 6, f"=PRODUCT(D{r},E{r})")
        bg, fg, bold = CLASE_FILL.get(clase, CLASE_FILL["normal"])
        ws.cell(r, 2).font = font(sz, bold, fg)
        if bg:
            ws.cell(r, 2).fill = fill(bg)
        for c in (3, 4, 5, 6):
            ws.cell(r, c).font = font(sz)
        ws.cell(r, 6).number_format = m.fmt_m
        ws.cell(r, 5).number_format = m.fmt_m
        obs = []
        etiqueta = {
            "solo_uno": "ROJO/DORADO: partida cotizada SOLO por {quien} (fuera de la base). Confirmar si se requiere.",
            "varios_fuera_base": "AZUL/DORADO: fuera de la base pero la cotizan varios postores: alcance real no considerado en la base.",
            "omitida": "CELESTE: partida no cotizada por {faltan} (se valoriza en la NIVELACIÓN).",
            "alerta": "ROJO SOBRE AMARILLO: ALERTA.",
            "referencial": "GRIS: referencial, no se suma.",
        }
        priced = {pid: m.precio(it, pid) for pid in m.ids if m.precio(it, pid)}
        faltan = [pid for pid in m.ids if not m.cubierto(it, anc, pid)]
        if clase in etiqueta:
            obs.append(etiqueta[clase].format(quien=", ".join(m.P[p]["corto"] for p in priced),
                                              faltan=", ".join(m.P[p]["corto"] for p in faltan)))
        if len(priced) >= 2:
            parc = [v[2] for v in priced.values()]
            if max(parc) - min(parc) > m.umbral:
                hi = max(priced, key=lambda p: priced[p][2])
                lo = min(priced, key=lambda p: priced[p][2])
                obs.insert(0 if clase == "alto_impacto" else len(obs),
                           f"AMARILLO CLARO: ALTO IMPACTO, diferencia de {m.moneda} {max(parc) - min(parc):,.0f} "
                           f"entre {m.P[lo]['corto']} y {m.P[hi]['corto']}.")
        for pid in m.ids:
            if pid in anc_global:
                obs.append(f"{m.P[pid]['corto']} no la detalla: incluida en su precio global del capítulo.")
            inc = ((it.get("precios") or {}).get(pid) or {}).get("incluido")
            if inc and pid not in priced:
                obs.append(f"{m.P[pid]['corto']}: {inc}")
        base_p = priced.get(m.base)
        pu_min = min((v[1] * (1 - (m.desc(p) or 0)) for p, v in priced.items()), default=None)
        for pid, d in cols.items():
            p = (it.get("precios") or {}).get(pid) or {}
            pr = priced.get(pid)
            if not pr:
                if pid not in anc_global and not p.get("incluido") and priced and clase != "referencial":
                    others = [f"{L(cols[o]['tot'])}{r}" for o in priced if o != pid]
                    if others:
                        ctx["no_cot"][pid].append(f"MIN({','.join(others)})" if len(others) > 1 else others[0])
                if p.get("incluido"):
                    x = ws.cell(r, d["cant"], "INCL.")
                    x.font = font(sz, True)
                    x.fill = fill(YELLOW) if "?" in str(p["incluido"]) else PatternFill()
                    if "?" in str(p["incluido"]):
                        x.value = "INCL.?"
                continue
            cant, pu, _ = pr
            ref = ctx["refs"].get(pid, {})
            rr = p.get("ref")
            f_q = font(sz)
            if rr and ref:
                keys = rr if isinstance(rr, list) else [rr]
                sh = ref["sheet"].replace("'", "''")
                if all(k in ref["rows"] for k in keys):
                    qc = "+".join(f"'{sh}'!{ref['qcol']}{ref['rows'][k]}" for k in keys)
                    pc = "+".join(f"'{sh}'!{ref['pcol']}{ref['rows'][k]}" for k in keys)
                    ws.cell(r, d["cant"], f"={qc}")
                    if m.usd:  # la cotización está en US$
                        ws.cell(r, d["pu_usd"], f"=IFERROR(({pc})/{L(d['cant'])}{r},0)")
                        ws.cell(r, d["pu"], f"={L(d['pu_usd'])}{r}*$F$5")
                    else:
                        ws.cell(r, d["pu"], f"=IFERROR(({pc})/{L(d['cant'])}{r},0)")
                    ws.cell(r, d["parcial"], f"={L(d['cant'])}{r}*{L(d['pu'])}{r}")
                    if "pu_neto" in d:
                        ws.cell(r, d["pu_neto"], f"={L(d['pu'])}{r}*(1-{ctx['rates'][pid]})")
                        ws.cell(r, d["parcial_neto"], f"={L(d['pu_neto'])}{r}*{L(d['cant'])}{r}")
                else:
                    write_price(ws, m, r, d, pid, cant, p.get("pu") if p.get("pu_usd") is None else None,
                                p.get("pu_usd"), f_q, m.fmt_m, ctx["rates"].get(pid))
            else:
                write_price(ws, m, r, d, pid, cant, p.get("pu") if p.get("pu_usd") is None else None,
                            p.get("pu_usd"), f_q, m.fmt_m, ctx["rates"].get(pid))
            for k in d:
                if k in ("cant", "pu_usd", "pu", "parcial", "pu_neto", "parcial_neto"):
                    ws.cell(r, d[k]).number_format = FMT_QTY if k == "cant" else (m.fmt_usd if k == "pu_usd" else m.fmt_m)
                    ws.cell(r, d[k]).font = font(sz)
            # cantidad azul: difiere > 2% de la base
            if base_p and pid != m.base and base_p[0]:
                dif = cant / base_p[0] - 1
                if abs(dif) > 0.02:
                    ws.cell(r, d["cant"]).font = font(sz, False, BLUE_TXT)
                    obs.append(f"AZUL (cant. {m.P[pid]['corto']}): metrado {cant:,.2f} vs base {base_p[0]:,.2f} ({dif:+.0%}).")
            pu_net = pu * (1 - (m.desc(pid) or 0))
            if pu_min and len(priced) >= 2 and pu_net >= 1.3 * pu_min:
                ws.cell(r, d["pu"]).font = font(sz, True, RED_TXT)
                obs.append(f"ROJO NEGRITA (P.U. {m.P[pid]['corto']}): {m.moneda} {pu_net:,.2f} vs {m.moneda} {pu_min:,.2f} "
                           f"({pu_net / pu_min - 1:+.0%}).")
            if p.get("excluir_subtotal"):
                obs.append(f"{m.P[pid]['corto']}: línea no sumada en su subtotal declarado (ver comentario en la Hoja 1).")
        if it.get("obs"):
            obs.append(it["obs"])
        if obs:
            _obs(ws, r, obs_col, "\n".join(obs))
        r += 1
    return rows, r


# ---------------------------------------------------------------- leyenda
LEYENDA = [
    ("Azul marino, letra blanca", NAVY, WHITE, False, "Columnas D-F, títulos de bloque", "Presupuesto meta (referencia). Si está vacío, falta el ppto. meta."),
    ("Verde intenso, letra blanca", GREEN, WHITE, True, "Encabezado de postor (Hoja 2)", "Postor recomendado (mejor relación precio / alcance / control)."),
    ("Verde limón", LIME, BLACK, True, "Encabezado de postor / DELTA", "Menor precio ofertado o bajo meta."),
    ("Amarillo", YELLOW, BLACK, True, "Encabezado de postor (Hoja 2)", "Postor en negociación u observado."),
    ("Naranja, letra roja", ORANGE, RED_TXT, True, "Encabezado de postor (Hoja 2)", "Oferta más cara / sobre meta."),
    ("Negro, letra amarilla", BLACK, YELLOW, True, "Encabezado de postor", "Versión anterior de una oferta."),
    ("Gris, negrita", GRAY, BLACK, True, "Descripción", "Precio global que cubre varias partidas: comparar por subtotal del capítulo."),
    ("Descripción roja sobre dorado claro", GOLD_LIGHT, RED_TXT, False, "Descripción", "Partida fuera de la base cotizada por UN solo postor. Decidir si se necesita."),
    ("Descripción azul sobre dorado claro", GOLD_LIGHT, BLUE_TXT, False, "Descripción", "Fuera de la base pero la cotizan VARIOS postores: alcance real no considerado."),
    ("Fondo celeste", CELESTE, BLACK, False, "Descripción", "Partida de la base que algún postor no cotiza ni declara incluida. Se valoriza en la NIVELACIÓN."),
    ("Fondo amarillo claro", YELLOW_LIGHT, BLACK, False, "Descripción", "Alto impacto: diferencia mayor al umbral entre postores."),
    ("Rojo sobre amarillo intenso", YELLOW, RED_TXT, True, "Descripción o P.U.", "Alerta crítica: total que no cuadra, duplicado, unidades distintas, metrado sin precio."),
    ("Celda amarilla \"INCL.?\"", YELLOW, BLACK, True, "Cantidad", "El postor no la separa; se asume incluida. Confirmar."),
    ("Cantidad en azul", None, BLUE_TXT, False, "Columna Cant.", "Metrado del postor difiere más de 2% del de la base."),
    ("P.U. rojo negrita", None, RED_TXT, True, "Columna P.U.", "P.U. 30% o más sobre el menor de los postores. Punto a negociar."),
    ("Número azul", None, BLUE_TXT, True, "GG / Util. / % dscto / TC", "Dato editable: simular descuentos o cambiar supuestos."),
    ("Texto azul negrita (criterios)", None, BLUE_TXT, True, "Bloque de criterios", "Ventaja / cumple requerimiento."),
    ("Texto rojo negrita (criterios)", None, RED_TXT, True, "Bloque de criterios", "Desventaja / riesgo de costo adicional."),
    ("Texto negro negrita (criterios)", None, BLACK, True, "Bloque de criterios", "Dato faltante o a confirmar con el postor."),
]


def hoja_leyenda(wb, m, h2name, cols2, rows2):
    ws = wb.create_sheet("LEYENDA Y NOTAS")
    ws.sheet_view.showGridLines = False
    for k, v in {"A": 2, "B": 40, "C": 22, "D": 22, "E": 22, "F": 60}.items():
        ws.column_dimensions[k].width = v
    extra = len(m.ids) - 3
    if extra > 0:
        ws.column_dimensions["F"].width = 22
    ws["B2"] = f"LEYENDA Y NOTAS - {m.s['especialidad']['titulo']}"
    ws["B2"].font = font(14, True)
    ws["B3"] = "Estándar de colores del cuadro comparativo. Cada fila resaltada de la Hoja 2 explica su motivo en OBSERVACIONES."
    ws["B3"].font = font(10, False, None, True)
    ws["B5"] = "1. CÓDIGO DE COLORES"
    ws["B5"].font = font(12, True)
    for c, t in zip("BCDEF", ["Formato (muestra)", "Dónde se aplica", None, None, "Qué significa / qué hacer"]):
        ws[f"{c}6"] = t
        ws[f"{c}6"].font = font(10, True, WHITE)
        ws[f"{c}6"].fill = fill(NAVY)
    r = 7
    for txt, bg, fg, b, donde, sig in LEYENDA:
        x = ws.cell(r, 2, txt)
        x.font = font(10, b, fg)
        if bg:
            x.fill = fill(bg)
        ws.cell(r, 3, donde).font = font(10)
        ws.cell(r, 6, sig).font = font(10)
        ws.cell(r, 6).alignment = WRAP_TL
        r += 1
    r += 1
    ws.cell(r, 2, "2. RESULTADO (vinculado al cuadro desglosado)").font = font(12, True)
    r += 1
    head = ["Concepto"] + [m.P[p]["corto"] for p in m.ids]
    for i, t in enumerate(head):
        x = ws.cell(r, 2 + i, t)
        x.font = font(10, True, WHITE)
        x.fill = fill(NAVY)
    r += 1
    sh = h2name.replace("'", "''")
    for lab, key, nf, use_ini in (("Precio ofertado sin IGV", "sub", m.fmt_m, False),
                                  ("Precio nivelado por alcance sin IGV", "niv", m.fmt_m, True),
                                  ("Variación % nivelada vs. menor", "niv_var", FMT_PCT, True)):
        ws.cell(r, 2, lab).font = font(10, True)
        for i, pid in enumerate(m.ids):
            d = cols2[pid]
            col = L(d["ini"] if use_ini else d["tot"])
            x = ws.cell(r, 3 + i, f"='{sh}'!{col}{rows2[key]}")
            x.number_format = nf
            x.font = font(10)
        r += 1
    r += 1
    ws.cell(r, 2, "3. CONCLUSIONES DEL INGENIERO DE COSTOS").font = font(12, True)
    r += 1
    for t in m.s.get("conclusiones", []):
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        x = ws.cell(r, 2, f"• {t}")
        x.font = font(10)
        x.alignment = WRAP_TL
        ws.row_dimensions[r].height = max(15, lines_needed(t, 150) * 13 + 3)
        r += 1


# ---------------------------------------------------------------- cotizaciones
def hoja_cotizacion(wb, m, pid, used, base_dir):
    p = m.P[pid]
    cot = p.get("cotizacion") or {}
    name = safe_sheet(cot.get("hoja") or p["corto"].upper(), used)
    src = cot.get("archivo_excel")
    if src:
        path = src if os.path.isabs(src) else os.path.join(base_dir, src)
        return name, copiar_hoja(wb, path, cot.get("hoja_origen"), name, cot)
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    for k, v in {"A": 2, "B": 12, "C": 70, "D": 8, "E": 12, "F": 14, "G": 16, "H": 14}.items():
        ws.column_dimensions[k].width = v
    ws["B1"] = cot.get("titulo", f"{p['nombre']} - COTIZACIÓN (transcrita)")
    ws["B1"].font = font(12, True)
    r = 2
    for k, v in (cot.get("encabezado") or {}).items():
        ws.cell(r, 2, k).font = font(10, True)
        ws.cell(r, 3, v).font = font(10)
        r += 1
    r += 1
    for i, t in enumerate(["Ítem", "Descripción", "Und", "Metrado", f"P.U. ({m.moneda if not m.usd else 'US$'})",
                           f"Parcial ({m.moneda if not m.usd else 'US$'})", "Control E×F-G"]):
        x = ws.cell(r, 2 + i, t)
        x.font = font(10, True, WHITE)
        x.fill = fill(NAVY)
    r += 1
    rows = {}
    first = r
    for f in cot.get("filas", []):
        ws.cell(r, 2, f.get("item"))
        ws.cell(r, 3, f.get("desc"))
        ws.cell(r, 4, f.get("und"))
        titulo = f.get("cant") is None and f.get("parcial") is None
        if not titulo:
            ws.cell(r, 5, f.get("cant"))
            ws.cell(r, 6, f.get("pu"))
            parcial = f.get("parcial")
            ws.cell(r, 7, parcial if parcial is not None else f"=E{r}*F{r}")
            ws.cell(r, 8, f"=E{r}*F{r}-G{r}")
            for c in (5, 6, 7, 8):
                ws.cell(r, c).number_format = FMT_NUM
            rows[str(f.get("item"))] = r
        for c in range(2, 9):
            ws.cell(r, c).font = font(10, titulo)
        r += 1
    last = r - 1
    for lab, val in (("SUBTOTAL (suma de partidas)", f"=SUM(G{first}:G{last})"),
                     ("TOTAL DECLARADO SIN IGV" + (" (US$)" if m.usd else ""),
                      (p.get("total_declarado_sin_igv") / m.tc if m.usd and p.get("total_declarado_sin_igv") is not None
                       else p.get("total_declarado_sin_igv"))),
                     ("DIFERENCIA (declarado - suma)", f"=G{r + 1}-G{r}")):
        ws.cell(r, 3, lab).font = font(10, True)
        x = ws.cell(r, 7, val)
        x.number_format = FMT_NUM
        x.font = font(10, True)
        r += 1
    r += 1
    for t in cot.get("notas", []):
        ws.cell(r, 3, f"* {t}").font = font(10)
        ws.cell(r, 3).alignment = WRAP_TL
        r += 1
    return name, {"sheet": name, "rows": rows, "qcol": "E", "pcol": "G"}


def copiar_hoja(wb, path, hoja_origen, name, cot):
    """Copia valores, formatos y combinaciones de la hoja del Excel del postor."""
    src_wb = load_workbook(path)
    src = src_wb[hoja_origen] if hoja_origen else src_wb.worksheets[0]
    ws = wb.create_sheet(name)
    from copy import copy
    for row in src.iter_rows():
        for c in row:
            if c.__class__.__name__ == "MergedCell":
                continue
            v = c.value
            if isinstance(v, str) and v.startswith("=") and "!" in v:
                v = src_wb[src.title][c.coordinate].value  # referencia externa: se deja la fórmula tal cual
            n = ws.cell(c.row, c.column, v)
            if c.has_style:
                n.font, n.fill, n.border = copy(c.font), copy(c.fill), copy(c.border)
                n.number_format, n.alignment = c.number_format, copy(c.alignment)
    for rg in src.merged_cells.ranges:
        ws.merge_cells(str(rg))
    for k, dim in src.column_dimensions.items():
        ws.column_dimensions[k].width = dim.width
    rows = {}
    cfg = cot.get("mapa") or {}  # {"col_item":"B","col_cant":"E","col_parcial":"G"}
    ci, cq, cp = cfg.get("col_item", "B"), cfg.get("col_cant", "E"), cfg.get("col_parcial", "G")
    for r in range(1, ws.max_row + 1):
        v = ws[f"{ci}{r}"].value
        if v is not None:
            rows[str(v)] = r
    return {"sheet": name, "rows": rows, "qcol": cq, "pcol": cp}


# ---------------------------------------------------------------- main
def nombre_archivo(spec):
    d, _ = fecha_larga(spec["proyecto"].get("fecha"))
    rev = int(spec["proyecto"].get("revision", 0))
    esp = spec["especialidad"]["nombre"]
    return f"{d:%y%m%d}_{spec['proyecto']['codigo']}_CC_{esp}_Rev.{rev:02d}.xlsx"


def generar(spec, out_dir, base_dir="."):
    m = Modelo(spec)
    wb = Workbook()
    esp = spec["especialidad"]["codigo"]
    proy = spec["proyecto"]["codigo"]
    used = set()
    n1 = safe_sheet(f"CC_{esp}-{proy}", used)
    n2 = safe_sheet(f"CC_{esp}_{proy}", used)
    used.add("LEYENDA Y NOTAS")
    hoja1(wb, m, n1)
    # las hojas de cotización se crean primero (para enlazar la Hoja 2) y luego se ordenan
    refs = {}
    cot_names = []
    for pid in m.ids:
        nm, ref = hoja_cotizacion(wb, m, pid, used, base_dir)
        refs[pid] = ref
        cot_names.append(nm)
    _, cols2, rows2 = hoja2(wb, m, n2, refs)
    hoja_leyenda(wb, m, n2, cols2, rows2)
    order = [n1, n2, "LEYENDA Y NOTAS"] + cot_names
    wb._sheets = [wb[n] for n in order]
    wb.active = 0
    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, nombre_archivo(spec))
    wb.save(path)
    return path, m


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    src = sys.argv[1]
    with open(src, encoding="utf-8") as fh:
        spec = json.load(fh)
    out = sys.argv[2] if len(sys.argv) > 2 else "."
    path, m = generar(spec, out, os.path.dirname(os.path.abspath(src)))
    print(path)
    for pid in m.ids:
        t = m.tot[pid]
        dec = m.P[pid].get("total_declarado_sin_igv")
        ok = "" if dec is None else ("OK" if abs(t["sub"] - float(dec)) < 0.5 else f"NO CUADRA (declarado {float(dec):,.2f})")
        print(f"  {m.P[pid]['corto']}: subtotal sin IGV {t['sub']:,.2f} {ok}")
