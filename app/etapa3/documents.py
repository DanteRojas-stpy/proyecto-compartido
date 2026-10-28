"""Lectura completa y localizada. Los originales no se modifican."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import os
import uuid
import zipfile
from pathlib import Path

import pymupdf
import pytesseract
import xlrd
from openpyxl import load_workbook
from PIL import Image


class DocumentError(ValueError):
    """Archivo no legible o fuera de los límites declarados."""


NAMESPACE = uuid.UUID("52bf8df4-e062-44c8-8e3c-5be90c5c3d37")


def stable_id(value: str) -> str:
    return str(uuid.uuid5(NAMESPACE, value))


def describe(data: bytes, filename: str, mime: str, version_id: str) -> dict:
    if not data or len(data) > int(os.getenv("E3_MAX_FILE_BYTES", "50000000")):
        raise DocumentError("Archivo vacío o mayor que el máximo de 50 MB.")
    digest = hashlib.sha256(data).hexdigest()
    return {"documento_id": stable_id(f"{version_id}:{digest}"),
            "storage_file_id": stable_id(f"archivo:{digest}"), "sha256": digest,
            "nombre_original": Path(filename).name, "mime_type": mime,
            "tamano_bytes": len(data)}


def _fragment(text: str, location: dict, document_id: str, serial: int) -> list[dict]:
    size = int(os.getenv("E3_CHUNK_CHARS", "400"))
    overlap = int(os.getenv("E3_CHUNK_OVERLAP", "60"))
    if not 0 <= overlap < size:
        raise DocumentError("Configuración de fragmentos inválida.")
    chunks = []
    for start in range(0, len(text), size - overlap):
        value = text[start:start + size].strip()
        if value:
            loc = {**location, "bloque": serial, "caracter_inicio": start}
            chunks.append({"id": stable_id(f"{document_id}:{json.dumps(loc, sort_keys=True)}:{value}"),
                           "documento_id": document_id, "texto": value, "ubicacion": loc})
        if start + size >= len(text):
            break
    return chunks


def _ocr(image: Image.Image, page: int, document_id: str) -> tuple[list[dict], float]:
    result = pytesseract.image_to_data(image.convert("RGB"), lang="spa+eng",
                                     output_type=pytesseract.Output.DICT, timeout=90)
    blocks: dict[tuple, list[int]] = {}
    for i, text in enumerate(result["text"]):
        if text.strip() and float(result["conf"][i]) >= 0:
            key = (result["block_num"][i], result["par_num"][i], result["line_num"][i])
            blocks.setdefault(key, []).append(i)
    fragments, confidences = [], []
    for serial, indexes in enumerate(blocks.values()):
        text = " ".join(result["text"][i] for i in indexes)
        confidence = sum(float(result["conf"][i]) for i in indexes) / (100 * len(indexes))
        confidences.append(confidence)
        x = min(result["left"][i] for i in indexes)
        y = min(result["top"][i] for i in indexes)
        right = max(result["left"][i] + result["width"][i] for i in indexes)
        bottom = max(result["top"][i] + result["height"][i] for i in indexes)
        fragments.extend(_fragment(text, {"pagina": page, "origen": "ocr", "confianza": confidence,
                                          "bbox": [x, y, right, bottom]}, document_id, serial))
    if not fragments:
        raise DocumentError(f"No se pudo leer el contenido de la página {page}.")
    return fragments, sum(confidences) / len(confidences)


def _cell(value):
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float):
        if not math.isfinite(value):
            raise DocumentError("La hoja contiene un número no finito.")
        return value
    return value.isoformat() if hasattr(value, "isoformat") else str(value)


def extract(data: bytes, meta: dict) -> dict:
    """Incluye toda página/hoja; superar un límite es error, nunca recorte silencioso."""
    document_id = meta["documento_id"]
    ext = Path(meta["nombre_original"]).suffix.lower()
    fragments, pages, sheets, observations = [], [], [], []
    use_ocr, use_rag, cells = False, False, 0
    if data.startswith(b"%PDF"):
        use_rag = True
        try:
            pdf = pymupdf.open(stream=data, filetype="pdf")
            if pdf.needs_pass or len(pdf) > 300:
                raise DocumentError("PDF protegido o con más de 300 páginas.")
            for page_number, page in enumerate(pdf, 1):
                text = page.get_text()
                image_area = sum(pymupdf.Rect(x["bbox"]).get_area()
                                 for x in page.get_image_info())
                needs_ocr = len(text.strip()) < 40 or (image_area > page.rect.get_area() * .6 and len(text) < 100)
                if needs_ocr:
                    use_ocr = True
                    pix = page.get_pixmap(matrix=pymupdf.Matrix(2, 2))
                    found, confidence = _ocr(Image.open(io.BytesIO(pix.tobytes("png"))), page_number, document_id)
                    fragments.extend(found)
                    if confidence < .55:
                        observations.append({"tipo": "confianza_insuficiente", "descripcion": f"La página {page_number} no permite interpretar datos con confianza suficiente.",
                                             "severidad": "alta", "requiere_aclaracion": True, "documento_id": document_id})
                else:
                    for serial, block in enumerate(page.get_text("blocks")):
                        if len(block) > 6 and block[6] == 0:
                            fragments.extend(_fragment(block[4], {"pagina": page_number, "origen": "pdf_texto", "bbox": list(block[:4])}, document_id, serial))
                pages.append({"pagina": page_number, "procesada": True, "uso_ocr": needs_ocr})
            pdf.close()
        except DocumentError:
            raise
        except Exception as exc:
            raise DocumentError("PDF dañado o no interpretable.") from exc
    elif ext in {".xlsx", ".xls"}:
        if ext == ".xlsx":
            try:
                with zipfile.ZipFile(io.BytesIO(data)) as archive:
                    if sum(x.file_size for x in archive.infolist()) > 200_000_000:
                        raise DocumentError("El Excel descomprimido supera 200 MB.")
                formulas = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
                values = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
                for ws in formulas.worksheets:
                    row_count = 0
                    cached = values[ws.title]
                    for row_index, (row, values_row) in enumerate(zip(ws.iter_rows(), cached.iter_rows()), 1):
                        cells += len(row)
                        if cells > 1_000_000:
                            raise DocumentError("El Excel supera un millón de celdas; no se guardó una lectura parcial.")
                        populated = []
                        for cell, value_cell in zip(row, values_row):
                            if cell.value is None:
                                continue
                            record = {"celda": cell.coordinate, "valor": _cell(value_cell.value),
                                      "original": _cell(cell.value), "formato": cell.number_format}
                            if cell.data_type == "f":
                                record["formula"] = cell.value
                                if value_cell.value is None:
                                    observations.append({"tipo": "formula_sin_resultado", "descripcion": f"La fórmula de {ws.title}!{cell.coordinate} no tiene resultado guardado.",
                                                         "severidad": "media", "requiere_aclaracion": True, "documento_id": document_id})
                            populated.append(record)
                        if populated:
                            text = json.dumps(populated, ensure_ascii=False)
                            # Una fila no se divide para no romper la relación columna/valor.
                            fragments.append({"id": stable_id(f"{document_id}:{ws.title}:{row_index}"), "documento_id": document_id,
                                              "texto": text, "ubicacion": {"hoja": ws.title, "fila": row_index,
                                               "columnas": [x["celda"] for x in populated], "origen": "xlsx"}})
                            row_count += 1
                    sheets.append({"hoja": ws.title, "procesada": True, "filas_con_datos": row_count, "oculta": ws.sheet_state != "visible"})
                formulas.close()
                values.close()
            except DocumentError:
                raise
            except Exception as exc:
                raise DocumentError("No se pudo abrir el XLSX.") from exc
        else:
            try:
                book = xlrd.open_workbook(file_contents=data, on_demand=True)
                for ws in book.sheets():
                    for r in range(ws.nrows):
                        cells += ws.ncols
                        if cells > 1_000_000:
                            raise DocumentError("El XLS supera un millón de celdas.")
                        text = json.dumps([{"columna": c + 1, "valor": _cell(ws.cell_value(r, c)), "tipo_celda": ws.cell_type(r, c)}
                                           for c in range(ws.ncols)], ensure_ascii=False)
                        fragments.append({"id": stable_id(f"{document_id}:{ws.name}:{r+1}"), "documento_id": document_id,
                                          "texto": text, "ubicacion": {"hoja": ws.name, "fila": r + 1, "origen": "xls"}})
                    sheets.append({"hoja": ws.name, "procesada": True, "filas_con_datos": ws.nrows})
                book.release_resources()
            except DocumentError:
                raise
            except Exception as exc:
                raise DocumentError("No se pudo abrir el XLS.") from exc
    elif ext == ".csv":
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("cp1252")
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        for index, row in enumerate(csv.reader(io.StringIO(text), dialect), 1):
            cells += len(row)
            if cells > 1_000_000:
                raise DocumentError("El CSV supera un millón de celdas.")
            fragments.append({"id": stable_id(f"{document_id}:csv:{index}"), "documento_id": document_id,
                              "texto": json.dumps(row, ensure_ascii=False), "ubicacion": {"hoja": "CSV", "fila": index,
                                "columnas": list(range(1, len(row) + 1)), "origen": "csv"}})
        sheets.append({"hoja": "CSV", "procesada": True, "filas_con_datos": len(fragments)})
    elif ext in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".webp", ".bmp"}:
        use_ocr, use_rag = True, True
        try:
            image = Image.open(io.BytesIO(data))
            if image.n_frames > 300:
                raise DocumentError("La imagen supera 300 páginas.")
            for frame in range(image.n_frames):
                image.seek(frame)
                found, confidence = _ocr(image, frame + 1, document_id)
                fragments.extend(found)
                pages.append({"pagina": frame + 1, "procesada": True, "uso_ocr": True})
                if confidence < .55:
                    observations.append({"tipo": "confianza_insuficiente", "descripcion": "La imagen contiene texto de confianza insuficiente.",
                                         "severidad": "alta", "requiere_aclaracion": True, "documento_id": document_id})
        except DocumentError:
            raise
        except Exception as exc:
            raise DocumentError("Imagen dañada o no interpretable.") from exc
    else:
        raise DocumentError("Formato no admitido. Se admiten PDF, imágenes, XLSX, XLS y CSV.")
    if not fragments:
        raise DocumentError("El documento no tiene contenido utilizable.")
    return {**meta, "fragmentos": fragments, "paginas": pages, "hojas": sheets,
            "uso_ocr": use_ocr, "uso_rag": use_rag, "observaciones": observations}
