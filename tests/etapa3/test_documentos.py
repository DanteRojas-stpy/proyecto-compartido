import io
import shutil
import uuid

import pymupdf
import pytest
from openpyxl import Workbook

from app.etapa3.documents import DocumentError, describe, extract


def meta(data, name):
    return describe(data, name, "application/octet-stream", str(uuid.uuid4()))


def test_sha_y_deduplicacion_por_contenido():
    version = str(uuid.uuid4())
    a, b = describe(b"dato", "uno.pdf", "application/pdf", version), describe(b"dato", "dos.pdf", "application/pdf", version)
    assert a["sha256"] == b["sha256"]
    assert a["documento_id"] == b["documento_id"]
    assert a["nombre_original"] != b["nombre_original"]


def test_pdf_texto_cubre_todas_las_paginas():
    book = pymupdf.open()
    for i in range(3):
        book.new_page().insert_text((40, 80), f"Equipo de prueba pagina {i+1}. Cantidad 2 UND precio USD 100 total USD 200. " * 3)
    data = book.tobytes()
    result = extract(data, meta(data, "ejemplo.pdf"))
    assert len(result["paginas"]) == 3
    assert {f["ubicacion"]["pagina"] for f in result["fragmentos"]} == {1, 2, 3}
    assert result["uso_rag"] and not result["uso_ocr"]


@pytest.mark.skipif(not shutil.which("tesseract"), reason="OCR se verifica en el contenedor con Tesseract")
def test_pdf_escaneado_y_imagen_usan_ocr_sin_observacion_por_usarlo():
    original = pymupdf.open()
    page = original.new_page()
    page.insert_text((40, 80), "CAPACIDAD 750 KG", fontsize=24)
    page.insert_text((40, 130), "PRECIO USD 100 TOTAL USD 200", fontsize=20)
    png = page.get_pixmap(matrix=pymupdf.Matrix(2, 2)).tobytes("png")
    scanned = pymupdf.open()
    scanned.new_page().insert_image(page.rect, stream=png)
    for content, name in [(scanned.tobytes(), "escaneado.pdf"), (png, "foto.png")]:
        result = extract(content, meta(content, name))
        assert result["uso_ocr"] and result["uso_rag"]
        assert "750" in " ".join(x["texto"] for x in result["fragmentos"])
        assert not any(o["tipo"] == "uso_ocr" for o in result["observaciones"])


def test_xlsx_lee_hojas_ocultas_y_no_inventa_resultado_formula():
    book = Workbook()
    book.active.title = "Oferta"
    book.active.append(["Equipo", "Cantidad", "Precio"])
    book.active.append(["Equipo A", 2, 100])
    hidden = book.create_sheet("Condiciones")
    hidden.sheet_state = "hidden"
    hidden.append(["Vigencia", "30 dias"])
    hidden.append(["Total", "=2*100"])
    stream = io.BytesIO()
    book.save(stream)
    data = stream.getvalue()
    result = extract(data, meta(data, "oferta.xlsx"))
    assert {x["hoja"] for x in result["hojas"]} == {"Oferta", "Condiciones"}
    assert any(o["tipo"] == "formula_sin_resultado" for o in result["observaciones"])
    assert not result["uso_rag"]
    assert any('"valor": null' in f["texto"] and '"formula"' in f["texto"] for f in result["fragmentos"])


def test_csv_cubre_filas_y_columnas_y_conserva_original():
    data = 'Equipo;Cantidad;Precio\nA;2;"1.250,50"\nB;3;200\n'.encode()
    result = extract(data, meta(data, "oferta.csv"))
    assert len(result["fragmentos"]) == 3
    assert result["fragmentos"][-1]["ubicacion"]["columnas"] == [1, 2, 3]
    assert "1.250,50" in result["fragmentos"][1]["texto"]
    assert not result["uso_rag"]


def test_xls_lee_todas_las_hojas():
    xlwt = pytest.importorskip("xlwt")
    book = xlwt.Workbook()
    for name in ["Oferta", "Ficha"]:
        sheet = book.add_sheet(name)
        sheet.write(0, 0, "Equipo de prueba")
        sheet.write(0, 1, 750)
    stream = io.BytesIO()
    book.save(stream)
    data = stream.getvalue()
    result = extract(data, meta(data, "oferta.xls"))
    assert {s["hoja"] for s in result["hojas"]} == {"Oferta", "Ficha"}
    assert len(result["fragmentos"]) == 2


def test_archivo_danado_y_formato_no_admitido_no_completan():
    for data, name in [(b"%PDF roto", "roto.pdf"), (b"contenido", "otro.docx")]:
        with pytest.raises(DocumentError):
            extract(data, meta(data, name))
