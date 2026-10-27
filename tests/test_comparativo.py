import openpyxl
import pytest

from app.comparativo.generar_comparativo import (
    Cotizacion,
    calcular_referencia,
    generar_comparativo,
    leer_filas,
)


@pytest.fixture
def plantilla(tmp_path):
    """Plantilla sintética con la misma forma que la del proyecto: base y bloques de proveedor."""
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "PLANTILLA"
    hoja["F5"] = 0
    hoja["D13"] = 2
    hoja["E13"] = 141344.6
    hoja["F13"] = "=E13*D13"
    for col_cant, col_usd in [("H", "I"), ("L", "M")]:
        hoja[f"{col_usd}13"] = 0
        hoja[f"{chr(ord(col_usd) + 1)}13"] = f"={col_usd}13*$F$5"
        hoja[f"{col_cant}13"] = 0
    hoja["H13"] = 0
    ruta = tmp_path / "plantilla.xlsx"
    libro.save(ruta)
    return str(ruta)


def test_escribe_precio_usd_y_cantidad_y_conserva_formulas(plantilla, tmp_path):
    salida = str(tmp_path / "salida.xlsx")
    cot = Cotizacion("PROVEEDOR EJEMPLO", "USD", 36000.0, 2.0)

    resultado = generar_comparativo(plantilla, {"PROVEEDOR 1": cot}, salida, tipo_cambio=3.6)

    libro = openpyxl.load_workbook(salida)
    hoja = libro["PLANTILLA"]
    assert hoja["H13"].value == 2
    assert hoja["I13"].value == 36000.0
    assert hoja["J13"].value == "=I13*$F$5"
    assert hoja["F5"].value == 3.6
    assert [r["proveedor"] for r in resultado.escritos] == ["PROVEEDOR EJEMPLO"]


def test_cotizacion_en_soles_queda_pendiente_sin_tocar_formulas(plantilla, tmp_path):
    salida = str(tmp_path / "salida.xlsx")
    cot = Cotizacion("PROVEEDOR PEN", "PEN", 37790.0, 2.0)

    resultado = generar_comparativo(plantilla, {"PROVEEDOR 2": cot}, salida)

    hoja = openpyxl.load_workbook(salida)["PLANTILLA"]
    assert hoja["M13"].value == 0
    assert hoja["N13"].value == "=M13*$F$5"
    assert resultado.escritos == []
    assert any("moneda" in p.get("motivo", "") for p in resultado.pendientes)


def test_sin_tipo_de_cambio_no_se_inventa_valor(plantilla, tmp_path):
    salida = str(tmp_path / "salida.xlsx")
    resultado = generar_comparativo(plantilla, {}, salida, tipo_cambio=None)

    hoja = openpyxl.load_workbook(salida)["PLANTILLA"]
    assert hoja["F5"].value == 0
    assert any("Tipo de cambio" in p.get("motivo", "") for p in resultado.pendientes)


def test_bloque_desconocido_se_rechaza(plantilla, tmp_path):
    with pytest.raises(ValueError):
        generar_comparativo(plantilla, {"PROVEEDOR X": Cotizacion("x", "USD", 1.0, 1.0)},
                            str(tmp_path / "s.xlsx"))


def test_referencia_calcula_desviacion_contra_la_base(plantilla):
    cot = {"PROVEEDOR 1": Cotizacion("A", "USD", 36000.0, 2.0)}

    filas = calcular_referencia(plantilla, cot, tipo_cambio=3.6)

    # Base: 2 x 141.344,60 = 282.689,20. Oferta: 2 x 36.000 x 3,6 = 259.200.
    assert filas[0]["parcial_s"] == pytest.approx(259200.0)
    assert filas[0]["desviacion_vs_base_pct"] == pytest.approx((259200 - 282689.2) / 282689.2 * 100, rel=1e-3)


def test_lee_filas_de_cotizacion_extraida(tmp_path):
    ruta = tmp_path / "cot.xlsx"
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Datos cotizacion"
    hoja.append(["Proveedor", "Moneda", "Precio_unitario", "Cantidad", "Descripcion", "Observaciones"])
    hoja.append(["PROVEEDOR DE PRUEBA", "USD", 100.0, 2, "Ítem de prueba", ""])
    hoja.append(["", "USD", 5.0, 1, "sin proveedor", ""])
    libro.save(ruta)

    filas = leer_filas(str(ruta))

    assert len(filas) == 1
    assert filas[0].proveedor == "PROVEEDOR DE PRUEBA"
    assert filas[0].precio_unitario == 100.0
    assert filas[0].cantidad == 2.0
